"""Exporta Test Plan y anexos actuales a Word; requiere python-docx."""
import csv
import re
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

root = Path(__file__).resolve().parents[1]
docs = root / 'docs'
document = Document()
section = document.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(2), Cm(1.8)
section.left_margin = section.right_margin = Cm(2)
for style in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    document.styles[style].font.name = 'Arial'
    document.styles[style].font.color.rgb = RGBColor(0, 0, 0)
    for border in document.styles[style].element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
document.styles['Normal'].font.size = Pt(10.5)
document.styles['Normal'].paragraph_format.space_after = Pt(6)
document.styles['Normal'].paragraph_format.line_spacing = 1.12
document.styles['Title'].font.size = Pt(28)
for style, size in [('Heading 1', 16), ('Heading 2', 12.5), ('Heading 3', 11)]:
    document.styles[style].font.size = Pt(size)
    document.styles[style].font.bold = True
    document.styles[style].paragraph_format.space_before = Pt(14)
    document.styles[style].paragraph_format.space_after = Pt(7)
    document.styles[style].paragraph_format.keep_with_next = True

def inline(paragraph, text):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    for index, part in enumerate(re.split(r'(\*\*.*?\*\*|`.*?`)', text)):
        if part.startswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name, run.font.size = 'Consolas', Pt(9)
        else:
            paragraph.add_run(part)

def table(rows):
    count = len(rows[0])
    lengths = [sum(min(100, len(row[column])) for row in rows) / len(rows) for column in range(count)]
    weights = [max(10, size) ** .65 for size in lengths]
    widths = [17 * size / sum(weights) for size in weights]
    if 'Probabilidad' in rows[0]:
        widths = [5.2, 2.5, 2.1, 7.2]
    grid = document.add_table(rows=1, cols=count)
    grid.autofit = False
    for column, width in zip(grid.columns, widths):
        column.width = Cm(width)
    for index, values in enumerate(rows):
        cells = grid.rows[0].cells if index == 0 else grid.add_row().cells
        for position, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Cm(widths[position])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            margin = OxmlElement('w:tcMar')
            for name in ['top', 'left', 'bottom', 'right']:
                side = OxmlElement(f'w:{name}')
                side.set(qn('w:w'), '80' if name in ['top', 'bottom'] else '100')
                side.set(qn('w:type'), 'dxa')
                margin.append(side)
            properties.append(margin)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '234E5A' if index == 0 else ('F2F5F7' if index % 2 else 'FFFFFF'))
            properties.append(shade)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.05
            inline(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9)
                if index == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
        # Keep rows intact; headings repeat when a table spans pages.
        row_property = grid.rows[index]._tr.get_or_add_trPr()
        row_property.append(OxmlElement('w:cantSplit'))
        if index == 0:
            row_property.append(OxmlElement('w:tblHeader'))
    borders = OxmlElement('w:tblBorders')
    for name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{name}')
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            border.set(qn(f'w:{key}'), value)
        borders.append(border)
    grid._tbl.tblPr.append(borders)
    document.add_paragraph().paragraph_format.space_after = Pt(2)

def markdown(source, heading_shift=0, skip_title=True):
    lines = source.splitlines()
    index = 0
    code = False
    while index < len(lines):
        line = lines[index].strip()
        if line.startswith('```'):
            code = not code
            index += 1
            continue
        if not line:
            index += 1
            continue
        if line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                row = [value.strip() for value in lines[index].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+', value) for value in row):
                    rows.append(row)
                index += 1
            table(rows)
            continue
        heading = re.match(r'^(#{1,6}) (.+)$', line)
        if heading:
            level = len(heading[1])
            if not (skip_title and level == 1):
                document.add_heading(heading[2], min(3, max(1, level - 1 + heading_shift)))
        else:
            paragraph = document.add_paragraph()
            if code:
                paragraph.add_run(line).font.name = 'Consolas'
            else:
                inline(paragraph, line)
        index += 1

document.add_paragraph('UNIVERSIDAD DEL VALLE', style='Subtitle')
document.add_paragraph('Univalle Bolivia', style='Subtitle')
document.add_paragraph('Software Quality Assurance', style='Subtitle')
document.add_paragraph('\n')
document.add_paragraph('Test Plan SQA de StockFácil', style='Title')
document.add_paragraph('Actualización del plan en un entorno ágil con Kanban')
document.add_paragraph('\n')
for text in ['Andre Romero Caballero', 'Santiago Camacho Echalar',
             'Docente Rolando Lara', '27 de octubre de 2026', 'Versión 1.0']:
    document.add_paragraph(text)
document.add_paragraph('\n')
document.add_paragraph('El plan documenta el alcance, las pruebas y el control de calidad del inventario. '
                       'Integra resultados reales de PHPUnit, Selenium y Cypress y conserva los pendientes de revisión y aceptación humanas.')
document.add_page_break()
source = (docs / 'Test_Plan_SQA.md').read_text(encoding='utf-8')
body = source[source.index('## 1 Introducción'):source.index('## Anexo A')]
markdown(body)
document.add_heading('Anexo A Casos de prueba y escenarios BDD', 1)
markdown((docs / 'Casos_de_Prueba.md').read_text(encoding='utf-8'), heading_shift=1)
document.add_heading('Escenarios BDD vinculados', 2)
feature = (root / 'tests/Browser/inventario.feature').read_text(encoding='utf-8')
for line in feature.splitlines():
    if not line.strip() or line.lstrip().startswith('#'):
        continue
    paragraph = document.add_paragraph(line.strip())
    paragraph.paragraph_format.space_after = Pt(3)
    if line.strip().startswith(('Escenario:', '@')):
        paragraph.paragraph_format.keep_with_next = True
document.add_heading('Anexo B Resultados de pruebas', 1)
markdown((docs / 'Resultados_Pruebas.md').read_text(encoding='utf-8'), heading_shift=1)
document.add_heading('Anexo C Flujo Kanban y tareas', 1)
markdown((docs / 'Kanban_Guia.md').read_text(encoding='utf-8'), heading_shift=1)
with (docs / 'Kanban_Tareas.csv').open(encoding='utf-8-sig') as stream:
    tasks = list(csv.DictReader(stream))
table([['Tarea', 'Responsable propuesto', 'Estado local']] +
      [[task['Titulo'], task['Responsable'], task['Estado']] for task in tasks])
document.add_paragraph('Los títulos K-01 a K-11 son IDs locales preparados. '
                       'GitHub_Entrega.md contiene los pasos exactos para crear Issues, tablero, PR y verificar Actions. '
                       'Pull_Request_SQA.md conserva la descripción lista para revisión. '
                       'Repositorio https://github.com/AndreRCRO/StockFacil.')
document.add_heading('Anexo D Auditoría interna', 1)
markdown((docs / 'Auditoria_Interna_SQA.md').read_text(encoding='utf-8'), heading_shift=1)
document.add_heading('Fuentes', 1)
markdown(source.split('## Fuentes\n', 1)[1])
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('StockFácil SQA  |  ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(8)
document.core_properties.title = 'Test Plan SQA de StockFácil'
document.core_properties.subject = 'Plan de calidad con Kanban y resultados verificables'
document.core_properties.author = 'Andre Romero Caballero y Santiago Camacho Echalar'
document.core_properties.comments = 'Revisión humana y aceptación pendientes según el informe.'
destination = docs / 'Test_Plan_SQA.docx'
document.save(destination)
print(f'DOCX generado: {destination}; tablas: {len(document.tables)}')
